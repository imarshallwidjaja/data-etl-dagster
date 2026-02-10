# =============================================================================
# Complex Spreadsheet Splitter Op
# =============================================================================
# Splits a complex XLSX spreadsheet into per-sheet long-form Parquet outputs,
# registers each output as an intermediate artifact, and publishes child
# manifests for downstream tabular ingestion.
# =============================================================================

from dataclasses import dataclass
import re
import tempfile
from pathlib import Path
from typing import Any

from dagster import In, OpExecutionContext, Out, op
from openpyxl import load_workbook
import polars as pl

from libs.models.manifest import Manifest
from libs.s3_utils import parse_s3_path
from libs.spatial_utils import normalize_headers

from .intermediate_artifacts import register_intermediate_from_local_file


__all__ = [
    "split_complex_spreadsheet_op",
    "_compose_multirow_headers",
    "_find_anchor_row",
    "_melt_with_id_columns",
    "_prepare_sheet_specs",
]


@dataclass(frozen=True)
class SheetAnchorSpec:
    sheet_name: str
    sheet_slug: str
    anchor_row: int


def _slugify_sheet_name(sheet_name: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", sheet_name.strip().lower()).strip("_")
    return slug or "sheet"


def _is_empty_cell(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def _find_anchor_row(worksheet, anchor_text: str, anchor_match: str) -> int | None:
    anchor = anchor_text.strip()
    if not anchor:
        raise ValueError("anchor_text must be non-empty")

    for row_idx, row in enumerate(worksheet.iter_rows(values_only=True)):
        for cell in row:
            if _is_empty_cell(cell):
                continue

            cell_text = str(cell).strip()

            if anchor_match == "exact":
                if cell_text.casefold() == anchor.casefold():
                    return row_idx
            elif anchor_match == "contains":
                if anchor.casefold() in cell_text.casefold():
                    return row_idx
            elif anchor_match == "regex":
                if re.search(anchor, cell_text, flags=re.IGNORECASE):
                    return row_idx
            else:
                raise ValueError(
                    f"Unsupported anchor_match '{anchor_match}'. "
                    "Expected one of: exact, contains, regex"
                )

    return None


def _prepare_sheet_specs(
    *,
    workbook,
    anchor_text: str,
    anchor_match: str,
    sheet_names: list[str] | None,
) -> tuple[list[SheetAnchorSpec], list[str]]:
    available_names = list(workbook.sheetnames)

    if sheet_names:
        requested = [name for name in sheet_names]
        missing = [name for name in requested if name not in available_names]
        if missing:
            raise ValueError(f"Configured sheet_names not found in workbook: {missing}")
        candidate_names = requested
    else:
        candidate_names = available_names

    specs: list[SheetAnchorSpec] = []
    skipped: list[str] = []

    for sheet_name in candidate_names:
        worksheet = workbook[sheet_name]
        anchor_row = _find_anchor_row(
            worksheet,
            anchor_text=anchor_text,
            anchor_match=anchor_match,
        )
        if anchor_row is None:
            skipped.append(sheet_name)
            continue

        specs.append(
            SheetAnchorSpec(
                sheet_name=sheet_name,
                sheet_slug=_slugify_sheet_name(sheet_name),
                anchor_row=anchor_row,
            )
        )

    if not specs:
        raise ValueError(
            f"No sheets matched anchor '{anchor_text}' using mode '{anchor_match}'"
        )

    return specs, skipped


def _trim_trailing_empty(frame: pl.DataFrame) -> pl.DataFrame:
    if frame.height == 0 or frame.width == 0:
        return frame

    rows = frame.rows()

    last_row_idx = -1
    last_col_idx = -1
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            if not _is_empty_cell(value):
                if row_idx > last_row_idx:
                    last_row_idx = row_idx
                if col_idx > last_col_idx:
                    last_col_idx = col_idx

    if last_row_idx < 0 or last_col_idx < 0:
        return frame.head(0)

    trimmed_rows = [list(row[: last_col_idx + 1]) for row in rows[: last_row_idx + 1]]
    trimmed_columns = frame.columns[: last_col_idx + 1]
    return pl.DataFrame(trimmed_rows, schema=trimmed_columns, orient="row")


def _drop_fully_empty_rows(frame: pl.DataFrame) -> pl.DataFrame:
    if frame.height == 0 or frame.width == 0:
        return frame

    rows = frame.rows()
    filtered_rows = [
        list(row) for row in rows if any(not _is_empty_cell(value) for value in row)
    ]

    if not filtered_rows:
        return frame.head(0)

    return pl.DataFrame(filtered_rows, schema=frame.columns, orient="row")


def _compose_multirow_headers(header_block: pl.DataFrame) -> list[str]:
    if header_block.height == 0:
        raise ValueError("Cannot compose headers from empty header block")

    rows = header_block.rows()
    raw_headers: list[str] = []

    for col_idx in range(header_block.width):
        parts: list[str] = []
        for row in rows:
            value = row[col_idx]
            if _is_empty_cell(value):
                continue
            parts.append(str(value).strip())

        if parts:
            raw_headers.append("_".join(parts))
        else:
            raw_headers.append(f"column_{col_idx + 1}")

    _, cleaned_headers = normalize_headers(raw_headers)
    return cleaned_headers


def _melt_with_id_columns(table: pl.DataFrame, id_column_count: int) -> pl.DataFrame:
    if id_column_count < 1:
        raise ValueError("id_column_count must be >= 1")
    if table.width <= id_column_count:
        raise ValueError(
            "id_column_count must be less than the number of table columns"
        )

    id_columns = table.columns[:id_column_count]
    value_columns = table.columns[id_column_count:]

    return table.unpivot(
        index=id_columns,
        on=value_columns,
        variable_name="variable",
        value_name="value",
    )


def _build_child_manifest(
    *,
    parent_manifest: Manifest,
    child_batch_id: str,
    child_dataset_id: str,
    sheet_name: str,
    blob_s3_path: str,
) -> dict[str, Any]:
    child_tags = parent_manifest.metadata.tags.copy()
    child_tags["dataset_id"] = child_dataset_id
    child_tags["parent_batch_id"] = parent_manifest.batch_id
    child_tags["source_sheet"] = sheet_name

    child_manifest_json = {
        "batch_id": child_batch_id,
        "uploader": parent_manifest.uploader,
        "intent": "ingest_tabular",
        "files": [
            {
                "path": blob_s3_path,
                "type": "tabular",
                "format": "Parquet",
            }
        ],
        "metadata": {
            "title": parent_manifest.metadata.title,
            "description": parent_manifest.metadata.description,
            "keywords": parent_manifest.metadata.keywords,
            "source": parent_manifest.metadata.source,
            "license": parent_manifest.metadata.license,
            "attribution": parent_manifest.metadata.attribution,
            "project": parent_manifest.metadata.project,
            "tags": child_tags,
            "join_config": None,
        },
    }

    return Manifest.model_validate(child_manifest_json).model_dump(mode="json")


def _download_spreadsheet_to_local_temp(*, minio, source_s3_path: str, log) -> str:
    bucket, key = parse_s3_path(source_s3_path)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    local_path = tmp.name
    tmp.close()

    try:
        if bucket == minio.landing_bucket:
            minio.download_from_landing(key, local_path)
        elif bucket == minio.lake_bucket:
            minio.download_from_lake(key, local_path)
        else:
            raise ValueError(
                f"Unsupported bucket '{bucket}' for spreadsheet download. "
                f"Expected '{minio.landing_bucket}' or '{minio.lake_bucket}'."
            )
        return local_path
    except Exception:
        Path(local_path).unlink(missing_ok=True)
        log.exception("Failed to download source spreadsheet")
        raise


def _process_sheet_to_long_table(
    *,
    local_xlsx_path: str,
    sheet_name: str,
    anchor_row: int,
    header_rows: int,
    id_column_count: int,
) -> pl.DataFrame:
    sheet_frame = pl.read_excel(
        source=local_xlsx_path,
        sheet_name=sheet_name,
        engine="calamine",
        has_header=False,
        drop_empty_rows=False,
        drop_empty_cols=False,
        raise_if_empty=False,
    )

    if sheet_frame is None or sheet_frame.height == 0:
        raise ValueError(f"Sheet '{sheet_name}' is empty")

    sliced = sheet_frame.slice(anchor_row)
    trimmed = _trim_trailing_empty(sliced)

    if trimmed.height < header_rows:
        raise ValueError(
            f"Sheet '{sheet_name}' does not have enough rows from anchor for header_rows={header_rows}"
        )

    header_block = trimmed.slice(0, header_rows)
    data_block = trimmed.slice(header_rows)
    data_block = _drop_fully_empty_rows(data_block)

    headers = _compose_multirow_headers(header_block)
    if len(headers) != data_block.width:
        raise ValueError(
            f"Header width/data width mismatch on sheet '{sheet_name}': "
            f"{len(headers)} != {data_block.width}"
        )

    renamed = data_block.rename(
        {old: new for old, new in zip(data_block.columns, headers)}
    )

    return _melt_with_id_columns(renamed, id_column_count=id_column_count)


def _split_complex_spreadsheet(
    *,
    minio,
    mongodb,
    manifest: dict[str, Any],
    dagster_run_id: str,
    log,
) -> dict[str, Any]:
    validated_manifest = Manifest(**manifest)
    if validated_manifest.intent != "ingest_complex_spreadsheet":
        raise ValueError(
            "split_complex_spreadsheet_op only supports intent='ingest_complex_spreadsheet'"
        )

    config = getattr(validated_manifest.metadata, "complex_spreadsheet", None)
    if config is None:
        raise ValueError("Manifest missing metadata.complex_spreadsheet configuration")

    source_s3_path = validated_manifest.files[0].path
    local_xlsx_path = _download_spreadsheet_to_local_temp(
        minio=minio,
        source_s3_path=source_s3_path,
        log=log,
    )

    workbook = None
    try:
        workbook = load_workbook(local_xlsx_path, read_only=True, data_only=True)

        sheet_specs, skipped_sheets = _prepare_sheet_specs(
            workbook=workbook,
            anchor_text=config.template_params.anchor_text,
            anchor_match=config.template_params.anchor_match,
            sheet_names=config.template_params.sheet_names,
        )

        base_dataset_id = str(validated_manifest.metadata.tags["dataset_id"]).strip()

        child_specs: list[dict[str, Any]] = []
        seen_slugs: set[str] = set()
        for spec in sheet_specs:
            if spec.sheet_slug in seen_slugs:
                raise ValueError(
                    f"Duplicate derived sheet slug '{spec.sheet_slug}' from workbook sheet names"
                )
            seen_slugs.add(spec.sheet_slug)

            child_batch_id = f"{validated_manifest.batch_id}__{spec.sheet_slug}"
            child_dataset_id = f"{base_dataset_id}__{spec.sheet_slug}"
            child_manifest_key = f"manifests/{child_batch_id}.json"

            child_specs.append(
                {
                    "sheet_name": spec.sheet_name,
                    "sheet_slug": spec.sheet_slug,
                    "anchor_row": spec.anchor_row,
                    "child_batch_id": child_batch_id,
                    "child_dataset_id": child_dataset_id,
                    "child_manifest_key": child_manifest_key,
                }
            )

        # Preflight: collision check for all child manifest keys before any side effects.
        for child in child_specs:
            if minio.object_exists_in_landing(child["child_manifest_key"]):
                raise FileExistsError(
                    f"Child manifest key already exists: {child['child_manifest_key']}"
                )

        run_id = mongodb.get_run_object_id(dagster_run_id)
        published_children: list[dict[str, Any]] = []

        for child in child_specs:
            long_table = _process_sheet_to_long_table(
                local_xlsx_path=local_xlsx_path,
                sheet_name=child["sheet_name"],
                anchor_row=child["anchor_row"],
                header_rows=config.template_params.header_rows,
                id_column_count=config.template_params.id_column_count,
            )

            tmp_parquet = tempfile.NamedTemporaryFile(delete=False, suffix=".parquet")
            tmp_parquet_path = tmp_parquet.name
            tmp_parquet.close()

            try:
                long_table.write_parquet(tmp_parquet_path)

                artifact_result = register_intermediate_from_local_file(
                    local_path=tmp_parquet_path,
                    batch_id=validated_manifest.batch_id,
                    run_id=run_id,
                    producer="split_complex_spreadsheet_op",
                    label=f"split_{child['sheet_slug']}",
                    parameters={
                        "template_id": config.template_id,
                        "sheet_name": child["sheet_name"],
                        "anchor_row": child["anchor_row"],
                        "header_rows": config.template_params.header_rows,
                        "id_column_count": config.template_params.id_column_count,
                    },
                    content_type="application/vnd.apache.parquet",
                    minio=minio,
                    mongodb=mongodb,
                    log=log,
                )
            finally:
                Path(tmp_parquet_path).unlink(missing_ok=True)

            child_manifest = _build_child_manifest(
                parent_manifest=validated_manifest,
                child_batch_id=child["child_batch_id"],
                child_dataset_id=child["child_dataset_id"],
                sheet_name=child["sheet_name"],
                blob_s3_path=artifact_result["blob_s3_path"],
            )

            minio.upload_json_to_landing(
                key=child["child_manifest_key"],
                payload=child_manifest,
                if_not_exists=True,
            )

            published_children.append(
                {
                    "sheet_name": child["sheet_name"],
                    "sheet_slug": child["sheet_slug"],
                    "child_batch_id": child["child_batch_id"],
                    "child_dataset_id": child["child_dataset_id"],
                    "child_manifest_key": child["child_manifest_key"],
                    "artifact_id": artifact_result["artifact_id"],
                    "blob_s3_path": artifact_result["blob_s3_path"],
                }
            )

        return {
            "manifest": validated_manifest.model_dump(mode="json"),
            "published_children": published_children,
            "skipped_sheets": skipped_sheets,
            "child_count": len(published_children),
        }
    finally:
        if workbook is not None:
            workbook.close()
        Path(local_xlsx_path).unlink(missing_ok=True)


@op(
    ins={"manifest": In(dagster_type=dict)},
    out={"split_result": Out(dagster_type=dict)},
    required_resource_keys={"minio", "mongodb"},
)
def split_complex_spreadsheet_op(context: OpExecutionContext, manifest: dict) -> dict:
    """Split complex spreadsheet into child tabular manifests."""
    return _split_complex_spreadsheet(
        minio=context.resources.minio,
        mongodb=context.resources.mongodb,
        manifest=manifest,
        dagster_run_id=context.run_id,
        log=context.log,
    )
